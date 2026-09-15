"""Flask web UI for the Literature Screening Pipeline (v2 — Master + Weekly merge)."""

import os
import re
import time
import tempfile
from pathlib import Path
from datetime import date

import pandas as pd
from flask import Flask, render_template, request, jsonify, send_from_directory

os.environ.setdefault("UNPAYWALL_EMAIL", "your.email@example.com")
import literature_pipeline as pipeline

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024


@app.route("/")
def index():
    return render_template("index.html")


def _normalize_doi(val):
    doi = str(val).strip().lower() if pd.notna(val) else ""
    if not doi:
        return ""
    for pfx in ("https://doi.org/", "http://doi.org/", "https://dx.doi.org/", "http://dx.doi.org/"):
        doi = doi.replace(pfx, "")
    doi = re.sub(r"^\s*doi\s*:\s*", "", doi)
    return doi.strip()


def _normalize_title(val):
    t = str(val).strip().lower() if pd.notna(val) else ""
    t = re.sub(r"[^\w\s]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


@app.route("/upload", methods=["POST"])
def upload():
    # --- Validate ---
    if "weekly" not in request.files:
        return jsonify({"error": "No weekly input file uploaded"}), 400
    weekly_file = request.files["weekly"]
    if not weekly_file.filename:
        return jsonify({"error": "No weekly file selected"}), 400

    email = request.form.get("email", "").strip()
    if not email:
        return jsonify({"error": "Email is required"}), 400

    batch_name = request.form.get("batch", "").strip() or f"Batch {date.today().isoformat()}"
    today = date.today().isoformat()

    # Configure pipeline
    pipeline.EMAIL = email
    pipeline.USER_AGENT = f"LiteratureScreeningTool/1.0 (mailto:{email})"
    pipeline.session.headers.update({"User-Agent": pipeline.USER_AGENT})

    # --- Parse files ---
    try:
        weekly_df = _read_upload(weekly_file)
        required = {"UI", "Title", "DOI"}
        missing = required - set(weekly_df.columns)
        if missing:
            return jsonify({"error": f"Weekly file missing columns: {', '.join(sorted(missing))}"}), 400

        master_df = pd.DataFrame()
        if "master" in request.files and request.files["master"].filename:
            master_df = _read_upload(request.files["master"])
    except Exception as exc:
        return jsonify({"error": f"File parsing error: {exc}"}), 400

    # --- Merge & detect duplicates ---
    try:
        result = _merge_and_detect(master_df, weekly_df, batch_name, today, email)
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500

    # --- Write output files ---
    try:
        output_dir = Path("output")
        output_dir.mkdir(parents=True, exist_ok=True)

        _write_updated_master(result["allRecords"], output_dir / "Updated_Master.xlsx")
        _write_screening(result["allRecords"], output_dir / "Screening_Output.xlsx")
        _write_log(result["newEntries"], output_dir / "Download_Log.xlsx")
        _write_purchase(result["purchase"], output_dir / "Purchase_List.xlsx")
    except Exception as exc:
        return jsonify({"error": f"Output generation error: {exc}"}), 500

    return jsonify(result)


def _read_upload(file_storage):
    ext = file_storage.filename.rsplit(".", 1)[-1].lower()
    with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
        file_storage.save(tmp.name)
        if ext == "csv":
            df = pd.read_csv(tmp.name)
        else:
            df = pd.read_excel(tmp.name)
        os.unlink(tmp.name)
    df.columns = [str(c).strip() for c in df.columns]
    for col in df.columns:
        df[col] = df[col].apply(lambda v: str(v).strip() if pd.notna(v) else "")
    return df


def _merge_and_detect(master_df, weekly_df, batch_name, today, email):
    # Build master lookup
    master_doi_set = set()
    master_title_set = set()
    master_records = []

    if not master_df.empty and "DOI" in master_df.columns:
        for _, row in master_df.iterrows():
            rec = row.to_dict()
            nd = _normalize_doi(rec.get("DOI", ""))
            nt = _normalize_title(rec.get("Title", ""))
            if nd:
                master_doi_set.add(nd)
            if nt:
                master_title_set.add(nt)
            rec["_source"] = "Master"
            rec["_rowType"] = "master"
            rec["_skipProcessing"] = True
            master_records.append(rec)

    # Process weekly records
    weekly_norm = []
    for _, row in weekly_df.iterrows():
        rec = row.to_dict()
        rec["_ndoi"] = _normalize_doi(rec.get("DOI", ""))
        rec["_ntitle"] = _normalize_title(rec.get("Title", ""))
        rec["_source"] = "New"
        rec["Date Added"] = today
        rec["Batch"] = batch_name
        weekly_norm.append(rec)

    # Count batch occurrences
    b_doi_cnt, b_title_cnt = {}, {}
    for r in weekly_norm:
        if r["_ndoi"]:
            b_doi_cnt[r["_ndoi"]] = b_doi_cnt.get(r["_ndoi"], 0) + 1
        if r["_ntitle"]:
            b_title_cnt[r["_ntitle"]] = b_title_cnt.get(r["_ntitle"], 0) + 1

    seen_dois, seen_titles = set(), set()
    master_dup_count = batch_dup_count = new_count = 0

    for r in weekly_norm:
        nd, nt = r["_ndoi"], r["_ntitle"]

        # Check vs master
        in_master_doi = nd and nd in master_doi_set
        in_master_title = nt and nt in master_title_set
        if in_master_doi or in_master_title:
            r["_rowType"] = "masterDup"
            r["_skipProcessing"] = True
            r["Status"] = "ALREADY_IN_MASTER"
            r["Notes"] = "Already exists in master file"
            r["Duplicate"] = True
            reasons = []
            if in_master_doi:
                reasons.append("DOI in master")
            if in_master_title:
                reasons.append("Title in master")
            r["Duplicate Reason"] = "; ".join(reasons)
            master_dup_count += 1
            continue

        # Check within batch
        b_doi_dup = nd and b_doi_cnt.get(nd, 0) > 1
        b_title_dup = nt and b_title_cnt.get(nt, 0) > 1
        if b_doi_dup or b_title_dup:
            r["_rowType"] = "batchDup"
            r["Duplicate"] = True
            reasons = []
            if b_doi_dup:
                reasons.append("Duplicate DOI in batch")
            if b_title_dup:
                reasons.append("Duplicate Title in batch")
            r["Duplicate Reason"] = "; ".join(reasons)

            first_doi = nd and nd not in seen_dois
            first_title = nt and nt not in seen_titles
            if first_doi or first_title:
                r["_skipProcessing"] = False
                if nd:
                    seen_dois.add(nd)
                if nt:
                    seen_titles.add(nt)
            else:
                r["_skipProcessing"] = True
                r["Status"] = "BATCH_DUPLICATE"
                r["Notes"] = "Duplicate within this batch"
            batch_dup_count += 1
            continue

        # Genuinely new
        if nd:
            seen_dois.add(nd)
        if nt:
            seen_titles.add(nt)
        r["_rowType"] = "new"
        r["_skipProcessing"] = False
        r["Duplicate"] = False
        r["Duplicate Reason"] = ""
        new_count += 1

    # Process new entries via APIs
    to_process = [r for r in weekly_norm if not r.get("_skipProcessing")]
    for r in to_process:
        doi = r["_ndoi"]
        try:
            result = pipeline.process_record({
                pipeline.UI_COLUMN: r.get("UI", ""),
                pipeline.TITLE_COLUMN: r.get("Title", ""),
                pipeline.DOI_COLUMN: doi,
                "Duplicate": r.get("Duplicate", False),
            })
            r["Status"] = result.get("Status", "")
            r["Publisher"] = result.get("Publisher", "")
            r["OA Source"] = result.get("OA Source", "")
            r["PDF URL"] = result.get("PDF URL", "")
            r["Publisher URL"] = result.get("Publisher URL", "")
            r["Purchase Cost"] = result.get("Purchase Cost", "")
            r["Notes"] = result.get("Notes", "")
        except Exception as exc:
            r["Status"] = "ERROR"
            r["Notes"] = str(exc)
        time.sleep(pipeline.REQUEST_DELAY)

    # Combine
    all_records = master_records + weekly_norm

    # Purchase list (new only)
    purchase = []
    for r in weekly_norm:
        if r.get("Status") == "PURCHASE_REQUIRED":
            purchase.append({
                "UI": r.get("UI", ""), "Title": r.get("Title", ""),
                "DOI": r.get("DOI", ""), "Publisher": r.get("Publisher", ""),
                "Website/site link for purchase": r.get("Publisher URL", ""),
                "Purchase Cost": r.get("Purchase Cost", ""),
            })

    stats = {
        "masterCount": len(master_records),
        "weeklyCount": len(weekly_norm),
        "masterDupCount": master_dup_count,
        "batchDupCount": batch_dup_count,
        "newCount": new_count,
        "toProcessCount": len(to_process),
        "totalCombined": len(all_records),
    }

    # Clean internal keys for JSON
    def clean(rec):
        out = {k: v for k, v in rec.items() if not k.startswith("_n")}
        return out

    return {
        "allRecords": [clean(r) for r in all_records],
        "newEntries": [clean(r) for r in weekly_norm],
        "purchase": purchase,
        "stats": stats,
    }


def _write_updated_master(records, path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    headers = ["UI", "Title", "DOI", "Status", "Publisher", "OA Source", "Date Added", "Batch"]
    ws.append(headers)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2D3142")
    for cell in ws[1]:
        cell.font = hfont
        cell.fill = hfill

    seen = set()
    for r in records:
        rt = r.get("_rowType", "")
        if rt == "masterDup":
            continue
        if rt == "batchDup" and r.get("_skipProcessing"):
            continue
        nd = _normalize_doi(r.get("DOI", ""))
        if nd and nd in seen:
            continue
        if nd:
            seen.add(nd)
        ws.append([r.get(h, "") for h in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([18, 80, 45, 22, 30, 30, 14, 20], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path)


def _write_screening(records, path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "All Records"
    headers = ["Source", "UI", "Title", "DOI", "Status", "Duplicate", "Duplicate Reason", "Batch"]
    ws.append(headers)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2D3142")
    for cell in ws[1]:
        cell.font = hfont
        cell.fill = hfill

    yellow = PatternFill("solid", fgColor="FFF2CC")
    orange = PatternFill("solid", fgColor="FFE0CC")
    gray = PatternFill("solid", fgColor="E8E8E8")

    for r in records:
        rt = r.get("_rowType", "")
        row_data = [
            "Master" if rt == "master" else "New",
            r.get("UI", ""), r.get("Title", ""), r.get("DOI", ""),
            r.get("Status", ""), "Yes" if r.get("Duplicate") else "No",
            r.get("Duplicate Reason", ""), r.get("Batch", ""),
        ]
        ws.append(row_data)
        fill = None
        if rt == "masterDup":
            fill = orange
        elif rt == "batchDup":
            fill = yellow
        elif rt == "master":
            fill = gray
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def _write_log(entries, path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "New Entries Log"
    headers = ["UI", "Title", "DOI", "Status", "OA Source", "Publisher", "Publisher URL", "Purchase Cost", "Notes"]
    ws.append(headers)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2D3142")
    for cell in ws[1]:
        cell.font = hfont
        cell.fill = hfill

    yellow = PatternFill("solid", fgColor="FFF2CC")
    orange = PatternFill("solid", fgColor="FFE0CC")

    for r in entries:
        ws.append([r.get(h, "") for h in headers])
        rt = r.get("_rowType", "")
        fill = orange if rt == "masterDup" else yellow if rt == "batchDup" else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    ws.freeze_panes = "A2"
    wb.save(path)


def _write_purchase(purchase, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase List"
    headers = ["UI", "Title", "DOI", "Publisher", "Website/site link for purchase", "Purchase Cost"]
    ws.append(headers)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2D3142")
    for cell in ws[1]:
        cell.font = hfont
        cell.fill = hfill

    for r in purchase:
        ws.append([r.get(h, "") for h in headers])

    ws.freeze_panes = "A2"
    wb.save(path)


@app.route("/download/<filename>")
def download(filename):
    allowed = {
        "Screening_Output.xlsx", "Download_Log.xlsx",
        "Purchase_List.xlsx", "Updated_Master.xlsx",
    }
    if filename not in allowed:
        return jsonify({"error": "File not found"}), 404
    output_dir = Path("output").resolve()
    fp = output_dir / filename
    if not fp.exists():
        return jsonify({"error": f"{filename} not generated yet"}), 404
    return send_from_directory(str(output_dir), filename, as_attachment=True)


if __name__ == "__main__":
    print("\n  📚 Literature Screening Pipeline — Web UI")
    print("  Open http://127.0.0.1:5000 in your browser\n")
    app.run(debug=True, host="127.0.0.1", port=5000)
