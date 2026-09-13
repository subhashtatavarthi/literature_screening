import os
import re
import time
import shutil
import zipfile
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import requests
import fitz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

INPUT_FILE = os.getenv("LIT_INPUT_FILE", "input/literature.xlsx")
INPUT_SHEET = os.getenv("LIT_INPUT_SHEET") or None
SCREENING_SHEET = "Titles for Screening"

OUTPUT_DIR = Path("output")
PDF_DIR = OUTPUT_DIR / "PDFs"
SCREENING_OUTPUT = OUTPUT_DIR / "Screening_Output.xlsx"
PURCHASE_OUTPUT = OUTPUT_DIR / "Purchase_List.xlsx"
LOG_OUTPUT = OUTPUT_DIR / "Download_Log.xlsx"
ZIP_OUTPUT = OUTPUT_DIR / "Literature_PDFs.zip"

EMAIL = os.getenv("UNPAYWALL_EMAIL", "your.email@example.com")
USER_AGENT = f"LiteratureScreeningTool/1.0 (mailto:{EMAIL})"
REQUEST_TIMEOUT = 30
REQUEST_DELAY = 0.5

UI_COLUMN = "UI"
TITLE_COLUMN = "Title"
DOI_COLUMN = "DOI"

session = requests.Session()
session.headers.update({"User-Agent": USER_AGENT, "Accept": "*/*"})


def clean_value(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize_ui(value):
    value = clean_value(value)
    return re.sub(r'[\\/:*?"<>|]', "_", value)


def normalize_doi(value):
    doi = clean_value(value).lower()
    if not doi:
        return ""
    for prefix in (
        "https://doi.org/", "http://doi.org/",
        "https://dx.doi.org/", "http://dx.doi.org/"
    ):
        doi = doi.replace(prefix, "")
    doi = re.sub(r"^\s*doi\s*:\s*", "", doi)
    return doi.strip()


def normalize_title(value):
    title = clean_value(value).lower()
    title = re.sub(r"[^\w\s]", " ", title)
    return re.sub(r"\s+", " ", title).strip()


def read_input_excel():
    df = pd.read_excel(INPUT_FILE, sheet_name=INPUT_SHEET) if INPUT_SHEET else pd.read_excel(INPUT_FILE)
    df.columns = [str(c).strip() for c in df.columns]
    required = {UI_COLUMN, TITLE_COLUMN, DOI_COLUMN}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    result = df[[UI_COLUMN, TITLE_COLUMN, DOI_COLUMN]].copy()
    for col in [UI_COLUMN, TITLE_COLUMN, DOI_COLUMN]:
        result[col] = result[col].apply(clean_value)
    return result


def detect_duplicates(df):
    df = df.copy()
    df["Normalized DOI"] = df[DOI_COLUMN].apply(normalize_doi)
    df["Normalized Title"] = df[TITLE_COLUMN].apply(normalize_title)
    doi_dup = df["Normalized DOI"].ne("") & df.duplicated("Normalized DOI", keep=False)
    title_dup = df["Normalized Title"].ne("") & df.duplicated("Normalized Title", keep=False)
    df["Duplicate"] = doi_dup | title_dup
    reasons = []
    for i in df.index:
        r = []
        if doi_dup.loc[i]:
            r.append("Duplicate DOI")
        if title_dup.loc[i]:
            r.append("Duplicate Title")
        reasons.append("; ".join(r))
    df["Duplicate Reason"] = reasons
    return df


def highlight_duplicates(filename):
    wb = load_workbook(filename)
    ws = wb[SCREENING_SHEET]
    fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    header = {cell.value: cell.column for cell in ws[1]}
    dup_col = header["Duplicate"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=dup_col).value is True:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A": 18, "B": 80, "C": 45, "D": 15, "E": 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wb.save(filename)


def create_screening_workbook(df):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(INPUT_FILE, SCREENING_OUTPUT)
    with pd.ExcelWriter(SCREENING_OUTPUT, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df[[UI_COLUMN, TITLE_COLUMN, DOI_COLUMN, "Duplicate", "Duplicate Reason"]].to_excel(
            writer, sheet_name=SCREENING_SHEET, index=False
        )
    highlight_duplicates(SCREENING_OUTPUT)


def crossref_lookup(doi):
    if not doi:
        return None
    try:
        r = session.get("https://api.crossref.org/works/" + quote(doi, safe=""), timeout=REQUEST_TIMEOUT)
        return r.json().get("message") if r.status_code == 200 else None
    except requests.RequestException:
        return None


def unpaywall_lookup(doi):
    if not doi:
        return None
    try:
        r = session.get(
            f"https://api.unpaywall.org/v2/{quote(doi, safe='')}",
            params={"email": EMAIL},
            timeout=REQUEST_TIMEOUT,
        )
        return r.json() if r.status_code == 200 else None
    except requests.RequestException:
        return None


def extract_oa_candidates(data):
    if not data:
        return []
    out = []
    best = data.get("best_oa_location")
    if best and best.get("url_for_pdf"):
        out.append({"url": best["url_for_pdf"], "source": "Unpaywall best OA location"})
    for loc in data.get("oa_locations") or []:
        url = loc.get("url_for_pdf")
        if url and not any(x["url"] == url for x in out):
            out.append({"url": url, "source": loc.get("host_type") or "Unpaywall OA location"})
    return out


def is_valid_pdf(path):
    try:
        if not path.exists() or path.stat().st_size < 1000:
            return False
        with open(path, "rb") as f:
            if f.read(4) != b"%PDF":
                return False
        doc = fitz.open(path)
        pages = doc.page_count
        doc.close()
        return pages > 0
    except Exception:
        return False


def download_pdf(url, destination):
    tmp = destination.with_suffix(".tmp")
    try:
        r = session.get(url, timeout=REQUEST_TIMEOUT, allow_redirects=True, stream=True)
        if r.status_code != 200:
            return False, r.url
        with open(tmp, "wb") as f:
            for chunk in r.iter_content(128 * 1024):
                if chunk:
                    f.write(chunk)
        if not is_valid_pdf(tmp):
            tmp.unlink(missing_ok=True)
            return False, r.url
        tmp.replace(destination)
        return True, r.url
    except Exception:
        tmp.unlink(missing_ok=True)
        return False, url


def process_record(row):
    ui = normalize_ui(row[UI_COLUMN])
    title = clean_value(row[TITLE_COLUMN])
    doi = normalize_doi(row[DOI_COLUMN])
    result = {
        "UI": ui, "Title": title, "DOI": doi, "Duplicate": bool(row["Duplicate"]),
        "Status": "", "OA Source": "", "PDF URL": "", "Publisher": "",
        "Publisher URL": "", "Purchase Cost": "", "Notes": ""
    }

    if not doi:
        result["Status"] = "MANUAL_REVIEW"
        result["Notes"] = "DOI missing"
        return result

    crossref = crossref_lookup(doi)
    if crossref:
        result["Publisher"] = crossref.get("publisher") or ""
        result["Publisher URL"] = crossref.get("URL") or f"https://doi.org/{doi}"
    else:
        result["Publisher URL"] = f"https://doi.org/{doi}"

    oa = unpaywall_lookup(doi)
    if not oa:
        result["Status"] = "MANUAL_REVIEW"
        result["Notes"] = "Unable to determine OA status"
        return result

    candidates = extract_oa_candidates(oa)
    if candidates:
        PDF_DIR.mkdir(parents=True, exist_ok=True)
        destination = PDF_DIR / f"{ui}.pdf"
        for c in candidates:
            ok, final_url = download_pdf(c["url"], destination)
            if ok:
                result.update({
                    "Status": "DOWNLOADED", "OA Source": c["source"],
                    "PDF URL": final_url, "Notes": f"Saved as {ui}.pdf"
                })
                return result
            time.sleep(REQUEST_DELAY)
        result.update({
            "Status": "OA_DOWNLOAD_FAILED", "OA Source": candidates[0]["source"],
            "PDF URL": candidates[0]["url"],
            "Notes": "Open-access location found, but automated PDF download failed."
        })
        return result

    if oa.get("is_oa") is False:
        result["Status"] = "PURCHASE_REQUIRED"
        result["Purchase Cost"] = "Price not publicly available"
        result["Notes"] = "No legitimate open-access location found."
        return result

    result["Status"] = "MANUAL_REVIEW"
    result["Notes"] = "No downloadable PDF URL was found."
    return result


def format_purchase_workbook():
    wb = load_workbook(PURCHASE_OUTPUT)
    ws = wb.active
    ws.title = "Purchase List"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {"A":18,"B":80,"C":45,"D":30,"E":70,"F":30}.items():
        ws.column_dimensions[col].width = width
    wb.save(PURCHASE_OUTPUT)


def create_purchase_list(results):
    purchase = results[results["Status"] == "PURCHASE_REQUIRED"].copy()
    purchase = purchase[["UI","Title","DOI","Publisher","Publisher URL","Purchase Cost"]]
    purchase.rename(columns={"Publisher URL": "Website/site link for purchase"}, inplace=True)
    purchase.to_excel(PURCHASE_OUTPUT, index=False)
    format_purchase_workbook()


def create_zip():
    ZIP_OUTPUT.unlink(missing_ok=True)
    with zipfile.ZipFile(ZIP_OUTPUT, "w", zipfile.ZIP_DEFLATED) as z:
        for pdf in PDF_DIR.glob("*.pdf"):
            z.write(pdf, arcname=pdf.name)


def main():
    if EMAIL == "your.email@example.com":
        raise ValueError(
            "Set UNPAYWALL_EMAIL to your email, e.g. "
            "UNPAYWALL_EMAIL=me@example.com python literature_pipeline.py"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)

    df = detect_duplicates(read_input_excel())
    create_screening_workbook(df)

    results = []
    total = len(df)
    for n, (_, row) in enumerate(df.iterrows(), 1):
        print(f"[{n}/{total}] Processing {row[UI_COLUMN]}")
        try:
            result = process_record(row)
        except Exception as exc:
            result = {
                "UI": normalize_ui(row[UI_COLUMN]), "Title": clean_value(row[TITLE_COLUMN]),
                "DOI": normalize_doi(row[DOI_COLUMN]), "Duplicate": bool(row["Duplicate"]),
                "Status": "ERROR", "OA Source": "", "PDF URL": "", "Publisher": "",
                "Publisher URL": "", "Purchase Cost": "", "Notes": str(exc)
            }
        results.append(result)
        print("   ", result["Status"])
        time.sleep(REQUEST_DELAY)

    results_df = pd.DataFrame(results)
    results_df.to_excel(LOG_OUTPUT, index=False)
    create_purchase_list(results_df)
    create_zip()

    print("\nComplete.")
    print(results_df["Status"].value_counts().to_string())
    print(f"\nOutputs are under: {OUTPUT_DIR.resolve()}")


if __name__ == "__main__":
    main()
